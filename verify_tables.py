"""Verify every reported number in the Chapter 5 tables against the evaluation outputs."""
import csv, re, os, sys, collections
from docx import Document
U=os.path.join(os.path.dirname(os.path.abspath(__file__)),'results')
DOC=sys.argv[1] if len(sys.argv)>1 else os.path.join(os.path.dirname(os.path.abspath(__file__)),'Dissertation.docx')
load=lambda n: list(csv.DictReader(open(os.path.join(U,n)))) if os.path.exists(os.path.join(U,n)) else None
num=lambda s: float(str(s).replace('−','-').replace('+','').replace('%','').strip())
d=Document(DOC)
ok=fail=0; missing=[]
def chk(lab,got,want,tol=0.006):
    global ok,fail
    g = got is not None and abs(num(got)-num(want))<=tol
    ok,fail=ok+g,fail+(not g)
    if not g: print(f"  FAIL {lab:56s} doc={want} file={got}")
    return g
def sect(n): print(f"\n{n}")

S={(r['gallery'],r['method']):r for r in load('summary.csv')}
C={(r['gallery'],r['comparison']):r for r in load('comparisons.csv') if r['metric']=='ap_e2e'}
A={(r['gallery'],r['branch']):r for r in load('alpha_results.csv')}
K=load('condition_results.csv'); R=load('d_resolution.csv')
RC={r['height']:r for r in load('recall_filtered.csv')}
TL=load('phase2_tiling_ablation.csv'); PC=load('d_per_candidate.csv')

sect("Table 5.1  detection recall, search-gallery column")
hmap={'below 64':'0-64','64 to 96':'64-96','96 to 128':'96-128','128 to 192':'128-192','192 to 256':'192-256','above 256':'256-+','all':'overall'}
for row in d.tables[12].rows[1:]:
    h=row.cells[0].text.strip(); k=hmap.get(h)
    r=RC.get(k)
    if not r: missing.append(f"5.1 row '{h}' (no matching row in recall_filtered.csv)"); continue
    chk(f"5.1 {h} search gallery", r['recall_pct'], row.cells[3].text.split('[')[0])

sect("Table 5.2  tiling")
first=list(TL[0].keys())[0]
TLD={r[first]:r for r in TL}
for row in d.tables[13].rows[1:]:
    h=row.cells[0].text.strip().replace(' to ','-').replace('above ','')
    if h=='below 64': h='32-64'   # the tiling ablation file labels this band 32-64
    r=TLD.get(h) or TLD.get(h+'+')
    if not r: missing.append(f"5.2 row '{h}'"); continue
    chk(f"5.2 {h} whole frame", r['full'], row.cells[1].text)
    chk(f"5.2 {h} tiled", r['tiled'], row.cells[2].text)

sect("Table 5.3  single-source retrieval (native)")
m53=[('Photograph, appearance model','OSNet photo x1'),('Photograph, vision-language model','IRRA photo x1'),
 ('Two photographs, appearance model','OSNet photos x2'),('Two photographs, vision-language mod','IRRA photos x2'),
 ('Two descriptions, averaged','Descriptions x2'),('Description','Description x1')]
for row in d.tables[14].rows[1:]:
    lab=row.cells[0].text.strip()
    key=next((v for k,v in m53 if lab.startswith(k[:30])),None)
    if not key: missing.append(f"5.3 row '{lab}'"); continue
    r=S[('native',key)]
    chk(f"5.3 {lab} e2e", r['ap_e2e'], row.cells[1].text.split('[')[0])
    chk(f"5.3 {lab} cond", r['ap_cond'], row.cells[2].text)
    chk(f"5.3 {lab} r1", r['r1_e2e'], row.cells[3].text)

sect("Table 5.4  resolution study")
RD={r['level']:r for r in R}
for row in d.tables[15].rows[1:]:
    lv=row.cells[0].text.strip().replace(' px','px')
    r=RD.get(lv)
    if not r: missing.append(f"5.4 level '{lv}'"); continue
    chk(f"5.4 {lv} appearance", r['osnet'], row.cells[1].text, 0.006)
    chk(f"5.4 {lv} vision-language", r['irra'], row.cells[2].text, 0.006)
    chk(f"5.4 {lv} description", r['text'], row.cells[3].text, 0.006)

sect("Table 5.5  tuned balance and description gain")
for row in d.tables[16].rows[1:]:
    gal=row.cells[0].text.strip().replace(" px","px")
    for br,base in (('osnet',1),('irra',5)):
        r=A[(gal,br)]
        chk(f"5.5 {gal}/{br} alpha", r['alpha_tuned'], row.cells[base].text)
        chk(f"5.5 {gal}/{br} photos", r['image_only'], row.cells[base+1].text)
        chk(f"5.5 {gal}/{br} fusion", r['fusion'], row.cells[base+2].text)
        chk(f"5.5 {gal}/{br} gain", r['desc_gain'], row.cells[base+3].text.replace('*','').split('[')[0])

sect("Table 5.6  quality-aware weighting")
PCD={r['branch']:r for r in PC}
row=d.tables[17].rows[1]
chk("5.6 per-candidate, vision-language", PCD['irra']['diff'], row.cells[2].text.replace('*','').split('[')[0])
import collections, random, statistics
CQ=list(csv.DictReader(open(os.path.join(U,'condition_perquery.csv'))))
def pooled(branch, baseline):
    dd=collections.defaultdict(dict)
    for r in CQ:
        if r['branch']!=branch or r['rerank']!='True': continue
        if r['rule'] in ('condition-aware',baseline): dd[(r['gallery'],r['qi'],r['pid'])][r['rule']]=float(r['ap_e2e'])
    per=collections.defaultdict(list)
    for (g,q,pid),v in dd.items():
        if len(v)==2: per[pid].append((v['condition-aware']-v[baseline])*100)
    return statistics.fmean([x for p in per for x in per[p]])
for br,col in (('osnet',1),('irra',2)):
    chk(f"5.6 per-condition vs native-tuned {br}", pooled(br,'fixed (native)'),
        d.tables[17].rows[2].cells[col].text.replace('*','').split('[')[0], 0.015)
    chk(f"5.6 per-condition vs best single {br}", pooled(br,'fixed (pooled)'),
        d.tables[17].rows[3].cells[col].text.replace('*','').split('[')[0], 0.015)
KD={(r['gallery'],r['branch'],r['baseline']):r for r in K if r['rerank']=='True'}
for rowi,base in ((4,'fixed (native)'),(5,'fixed (pooled)')):
    for br,col in (('osnet',1),('irra',2)):
        cell=d.tables[17].rows[rowi].cells[col].text.replace('*','').strip()
        r=KD[('16px',br,base)]
        lo,hi=re.search(r'\[([^,]+),\s*([^\]]+)\]',cell).groups()
        chk(f"5.6 16px {br} vs {base}", r['diff'], cell.split('[')[0])
        chk(f"5.6 16px {br} lo", r['lo'], lo); chk(f"5.6 16px {br} hi", r['hi'], hi)

sect("Table 5.9 / 5.10  systems and components")
m59={'Photograph, appearance':'OSNet photo x1','Photograph, vision-language':'IRRA photo x1','Description':'Description x1',
 'Two photographs, appearance':'OSNet photos x2','Two photographs, vision-language':'IRRA photos x2',
 'System, appearance model':'Recommended','System, vision-language model':'Recommended (IRRA image branch)'}
for row in d.tables[20].rows[1:]:
    lab=row.cells[0].text.strip()
    for gal,col in (('native',1),('mixed',2)):
        chk(f"5.9 {lab}/{gal}", S[(gal,m59[lab])]['ap_e2e'], row.cells[col].text.split('[')[0])
NM={'Remove fusion':'remove fusion','Remove second photograph':'remove second photograph','Remove second description':'remove second description',
 'Remove re-ranking':'remove re-ranking','Add query expansion':'add query expansion','Add resolution matching':'add resolution matching'}
for row in d.tables[21].rows[1:]:
    lab=row.cells[0].text.strip()
    for gal,col in (('native',1),('mixed',2)):
        cell=row.cells[col].text.strip(); r=C[(gal,NM[lab])]
        lo,hi=re.search(r'\[([^,]+),\s*([^\]]+)\]',cell).groups()
        chk(f"5.10 {lab}/{gal}", r['diff'], cell.split('[')[0])
        chk(f"5.10 {lab}/{gal} lo", r['lo'], lo); chk(f"5.10 {lab}/{gal} hi", r['hi'], hi)


sect("Sections 5.9.1 / 6.4  description contribution against a named baseline")
# Uses cboot() exactly as defined in step1_corrected_eval.py (numpy default_rng,
# seed 0, 5,000 identity-level resamples, np.percentile at 2.5 / 97.5), so every
# interval below is reproducible with the submitted code at its default settings.
import numpy as _np
def cboot(vals, pids, reps=5000, seed=0):
    vals=_np.asarray(vals,float); pids=_np.asarray(pids)
    ok=~_np.isnan(vals); v,p=vals[ok],pids[ok]
    u,inv=_np.unique(p,return_inverse=True)
    sums=_np.bincount(inv,weights=v); cnt=_np.bincount(inv).astype(float)
    rng=_np.random.default_rng(seed)
    idx=rng.integers(0,len(u),size=(reps,len(u)))
    means=sums[idx].sum(1)/cnt[idx].sum(1)
    return v.mean(), float(_np.percentile(means,2.5)), float(_np.percentile(means,97.5))
AQ=[x for x in csv.DictReader(open(os.path.join(U,'alpha_perquery.csv'))) if x['gallery']=='native' and x['branch']=='irra']
PQ=[x for x in csv.DictReader(open(os.path.join(U,'perquery.csv'))) if x['gallery']=='native']
Am={(x['method'],x['qi']):(float(x['ap_e2e'])*100,x['pid']) for x in AQ}
Pm={(x['method'],x['qi']):(float(x['ap_e2e'])*100,x['pid']) for x in PQ}
QS=sorted({k[1] for k in Pm})
agree=max(abs(Pm[('IRRA photos x2',q)][0]-Am[('image only',q)][0]) for q in QS)
chk("the two per-query files agree on the shared method", agree, 0.0, 1e-9)
base=sum(Am[('image only + re-ranking',q)][0] for q in QS)/len(QS)
chk("baseline: two photographs + re-ranking, no description", base, 60.86, 0.005)
def gain(af,bf,akey,bkey):
    d=[af[(akey,q)][0]-bf[(bkey,q)][0] for q in QS]
    pid=[af[(akey,q)][1] for q in QS]
    return cboot(d,pid)
CASES=[("shared balance 0.80 + re-ranking",Pm,Am,'Recommended (IRRA image branch)','image only + re-ranking',"+0.02","-0.57","+0.58"),
 ("re-tuned balance 0.95 + re-ranking",Am,Am,'fusion + re-ranking','image only + re-ranking',"+0.31","+0.11","+0.55"),
 ("shared balance 0.80, no re-ranking",Pm,Pm,'IRRA fusion, no re-ranking','IRRA photos x2',"+0.34","-0.17","+0.85"),
 ("re-tuned balance 0.95, no re-ranking",Am,Am,'fusion','image only',"+0.30","+0.14","+0.45"),
 ("re-tuned vs shared, the two configurations",Am,Pm,'fusion + re-ranking','Recommended (IRRA image branch)',"+0.29","-0.16","+0.76")]
for lab,af,bf,ak,bk,m,lo,hi in CASES:
    o,l,h=gain(af,bf,ak,bk)
    chk(f"description gain, {lab}", o, m, 0.006)
    chk(f"  its interval, {lab}", l, lo, 0.006); chk(f"  its interval, {lab} (hi)", h, hi, 0.006)
# both re-ranked comparisons must rest on the same no-description baseline
i1=sum(Pm[('Recommended (IRRA image branch)',q)][0] for q in QS)/len(QS)-gain(Pm,Am,'Recommended (IRRA image branch)','image only + re-ranking')[0]
i2=sum(Am[('fusion + re-ranking',q)][0] for q in QS)/len(QS)-gain(Am,Am,'fusion + re-ranking','image only + re-ranking')[0]
chk("implied baseline identical in both re-ranked comparisons", i1, f"{i2:.4f}", 0.005)


sect("Table 5.7  sensitivity to the description   /   Table 6.1  failure analysis")
import pandas as _pd
E=_pd.read_csv(os.path.join(U,'d_exploratory_perquery.csv'))
_dt=(E.ap_txt1-E.ap_txt2).abs()*100; _df=(E.ap_fus1-E.ap_fus2).abs()*100
T57=d.tables[18]
chk("5.7 mean difference, description alone", _dt.mean(), T57.rows[1].cells[1].text.split()[0])
chk("5.7 mean difference, fused", _df.mean(), T57.rows[2].cells[1].text.split()[0])
_m,_lo,_hi=cboot((_dt-_df).values,E.pid.values)
_c=T57.rows[3].cells[1].text; _l,_h=re.search(r'\[([^,]+),\s*([^\]]+)\]',_c).groups()
chk("5.7 reduction from combining", _m, _c.split('[')[0]); chk("5.7 reduction lo", _lo, _l); chk("5.7 reduction hi", _hi, _h)
chk("5.7 searches differing by >40 points", (_dt>40).sum(), T57.rows[4].cells[1].text.split('(')[1].split()[0])
chk("5.7 description-only, first", E.ap_txt1.mean()*100, T57.rows[5].cells[1].text.split(' to ')[0])
chk("5.7 description-only, averaged", E.ap_txt_avg.mean()*100, T57.rows[5].cells[1].text.split(' to ')[1])
_m,_lo,_hi=cboot(((E.ap_txt_avg-E.ap_txt1)*100).values,E.pid.values)
_c=T57.rows[6].cells[1].text; _l,_h=re.search(r'\[([^,]+),\s*([^\]]+)\]',_c).groups()
chk("5.7 gain from averaging", _m, _c.split('[')[0]); chk("5.7 averaging lo", _lo, _l); chk("5.7 averaging hi", _hi, _h)
PP=_pd.read_csv(os.path.join(U,'perquery.csv'))
PP=PP[(PP.gallery=='native')&(PP.method=='Recommended (IRRA image branch)')].copy()
PP['frac']=PP.n_detected/PP.n_gt
T61=d.tables[22]
SEL=[PP.frac<=0.75,(PP.frac>0.75)&(PP.frac<0.99),PP.frac>=0.99,PP.n_gt<=5,(PP.n_gt>=6)&(PP.n_gt<=10),(PP.n_gt>=11)&(PP.n_gt<=20),PP.n_gt>20]
for i,sel in enumerate(SEL):
    g=PP[sel]; row=T61.rows[i+1]
    chk(f"6.1 row {i+1} searches", len(g), row.cells[1].text)
    chk(f"6.1 row {i+1} Rank-1", g.r1_e2e.mean()*100, row.cells[2].text.replace('%',''), 0.06)
    chk(f"6.1 row {i+1} end-to-end mAP", g.ap_e2e.mean()*100, row.cells[3].text, 0.06)
    if row.cells[4].text.strip() not in ('—','-',''):
        chk(f"6.1 row {i+1} conditional mAP", g.ap_cond.mean()*100, row.cells[4].text, 0.06)

print(f"\n{'='*72}\nCHECKS PASSED: {ok}    FAILED: {fail}")
if missing:
    print("NOT CHECKED (no matching row in the output files):")
    for m in missing: print("   -", m)
for f in ['e_confidence.txt','e_confidence_perquery.csv']:
    if not os.path.exists(os.path.join(U,f)): print(f"MISSING OUTPUT FILE: {f}  (Table 5.8 / Figure 5.6)")
sys.exit(1 if fail else 0)
